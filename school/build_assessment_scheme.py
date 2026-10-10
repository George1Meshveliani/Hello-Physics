# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Arial"
ACCENT = RGBColor(0x2C, 0x6E, 0x6A)
DARK = RGBColor(0x1E, 0x24, 0x33)

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'BFBFBF')
        borders.append(el)
    tcPr.append(borders)

def style_run(run, size=11, bold=False, color=None, italic=False):
    run.font.name = FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT)
    rFonts.set(qn('w:cs'), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color

def add_para(doc, text, size=11, bold=False, color=None, space_after=6, space_before=0, align=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align:
        p.alignment = align
    r = p.add_run(text)
    style_run(r, size=size, bold=bold, color=color, italic=italic)
    return p

def add_heading(doc, text, level=1):
    if level == 1:
        p = add_para(doc, text, size=20, bold=True, color=ACCENT, space_after=4, space_before=18)
    else:
        p = add_para(doc, text, size=14, bold=True, color=DARK, space_after=6, space_before=14)
    return p

def make_table(doc, headers, rows, col_widths_cm, header_fill="2C6E6A"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].width = Cm(col_widths_cm[i])
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        style_run(r, size=10.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        set_cell_shading(hdr_cells[i], header_fill)
        set_cell_borders(hdr_cells[i])
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Cm(col_widths_cm[i])
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i > 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            style_run(r, size=10, bold=(i == 0))
            set_cell_borders(cells[i])
    return table

doc = Document()

section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.left_margin = Cm(2)
section.right_margin = Cm(2)
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)

style = doc.styles['Normal']
style.font.name = FONT
style.font.size = Pt(11)
rPr = style.element.get_or_add_rPr()
rFonts = rPr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = OxmlElement('w:rFonts')
    rPr.append(rFonts)
rFonts.set(qn('w:eastAsia'), FONT)
rFonts.set(qn('w:cs'), FONT)

# ---------- Title ----------
add_para(doc, "შეფასების სისტემა", size=26, bold=True, color=ACCENT, space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "ფიზიკა · VII კლასი — 10-ქულიანი შეფასების სქემა", size=13, bold=False,
         color=RGBColor(0x56, 0x5F, 0x75), space_after=18, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)

# ---------- 1. ზოგადი პრინციპები ----------
add_heading(doc, "1. ზოგადი პრინციპები")
add_para(doc, "მოსწავლის მიღწევები ფასდება 10-ქულიანი სკალით, ეროვნული სასწავლო გეგმის შეფასების პრინციპების "
              "შესაბამისად. გამოიყენება შეფასების ორი ურთიერთდამატებითი ფორმა:")
add_para(doc, "•  ფორმირებითი (მიმდინარე) შეფასება — მიმდინარეობს ყოველ გაკვეთილზე, აძლევს მოსწავლესა და "
              "მასწავლებელს უკუკავშირს სწავლის პროცესში და არ არის აუცილებლად დაკავშირებული ქულასთან.", space_after=3)
add_para(doc, "•  განმსაზღვრელი (საბოლოო) შეფასება — აჯამებს გარკვეულ პერიოდში მიღწეულ შედეგს (საკონტროლო "
              "სამუშაო, ლაბორატორიული ანგარიში, პროექტი, სემესტრული/წლიური გამოცდა) და აისახება 10-ქულიან სკალაზე.",
         space_after=10)
add_para(doc, "შეფასების ყველა ფორმა ეყრდნობა წინასწარ განსაზღვრულ და მოსწავლისთვის ცნობილ კრიტერიუმებს (რუბრიკებს), "
              "რათა შეფასება იყოს გამჭვირვალე და სამართლიანი.")

# ---------- 2. 10-ქულიანი სკალა ----------
add_heading(doc, "2. შეფასების 10-ქულიანი სკალა")
scale_headers = ["ქულა", "დონე", "აღწერა"]
scale_rows = [
    ["9–10", "შესანიშნავი", "მასალა აქვს სრულად და ღრმად გააზრებული; ამოცანებს/ცდებს ასრულებს დამოუკიდებლად, "
                            "ზუსტად და უშეცდომოდ; დასკვნებს სათანადოდ ასაბუთებს."],
    ["7–8", "კარგი", "მასალა ესმის კარგად; დავალებათა უმეტესობას წყვეტს სწორად, დაშვებულია უმნიშვნელო უზუსტობა."],
    ["5–6", "საშუალო", "ძირითადი ცნებები ესმის, მაგრამ გამოყენებაში/გამოთვლებში ხდება შეცდომები; ესაჭიროება "
                       "მასწავლებლის მცირედი დახმარება."],
    ["3–4", "დამაკმაყოფილებელი", "მასალის მხოლოდ ნაწილი აქვს გაგებული; დამოუკიდებლად ვერ წყვეტს დავალებას; "
                                 "მნიშვნელოვანი შეცდომებია."],
    ["1–2", "მინიმალური", "ცოდნა ფრაგმენტულია; დავალებას ასრულებს მხოლოდ ნაწილობრივ ან არასწორად."],
    ["0", "არადამაკმაყოფილებელი", "დავალება არ არის შესრულებული ან სრულად არასწორია."],
]
make_table(doc, scale_headers, scale_rows, [2.2, 4.3, 10.5])
doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ---------- 3. შეფასების ფორმები ----------
add_heading(doc, "3. შეფასების ფორმები ფიზიკაში (VII კლასი)")
add_para(doc, "ფორმირებითი შეფასება: ზეპირი გამოკითხვა, საკლასო აქტივობის დაკვირვება, საშინაო დავალება, "
              "წერითი სამუშაო ფურცელი.", space_after=3)
add_para(doc, "განმსაზღვრელი შეფასება: საკონტროლო სამუშაო, ლაბორატორიული ანგარიშის შეფასება რუბრიკით, "
              "პროექტის/პრეზენტაციის შეფასება, სემესტრული და წლიური შემაჯამებელი (საგამოცდო) წერითი სამუშაო.",
         space_after=10)

# ---------- 4. ზოგადი კრიტერიუმები ----------
add_heading(doc, "4. ზოგადი შეფასების კრიტერიუმები")
add_para(doc, "ქვემოთ მოცემული ოთხი კრიტერიუმი გამოიყენება შეფასების ნებისმიერი კონკრეტული ფორმისთვის "
              "(საკონტროლო, ლაბორატორიული, პროექტი და ა.შ.) რუბრიკის შედგენისას:", space_after=8)
crit_headers = ["კრიტერიუმი", "რას აფასებს"]
crit_rows = [
    ["ცოდნა და გაგება", "ფიზიკური ცნებების, კანონებისა და ფორმულების ცოდნა და გაგება."],
    ["გამოყენება", "ცოდნის გამოყენება ამოცანის ამოხსნისას ან ცდის/ლაბორატორიული სამუშაოს ჩატარებისას."],
    ["ანალიზი და დასკვნა", "მიღებული შედეგის ინტერპრეტაცია, დასაბუთება და დასკვნის ჩამოყალიბება."],
    ["კომუნიკაცია და გაფორმება", "პასუხის/ანგარიშის თანმიმდევრულობა, ფიზიკური ენისა და ერთეულების სწორად გამოყენება."],
]
make_table(doc, crit_headers, crit_rows, [5.5, 11.5])
doc.add_paragraph().paragraph_format.space_after = Pt(4)

# ---------- 5. მაგალითი ----------
add_heading(doc, "5. მაგალითი — VII კლასი")
add_heading(doc, "ლაბორატორიული სამუშაოს შეფასების რუბრიკა: „სხეულის სიმკვრივის განსაზღვრა“", level=2)
add_para(doc, "თემატური ბლოკი: მასა და სიმკვრივე (კვირა XXI). მოსწავლე განსაზღვრავს მყარი სხეულის სიმკვრივეს "
              "მასისა და მოცულობის გაზომვით და შედეგს ადარებს ცხრილურ მონაცემს.", size=10.5, italic=True,
         color=RGBColor(0x56, 0x5F, 0x75), space_after=10)

rubric_headers = ["კრიტერიუმი", "მაქს. ქულა", "შეფასების აღწერა"]
rubric_rows = [
    ["1. თეორიული საფუძველი და მიზანი", "2", "მოსწავლემ იცის სიმკვრივის ფორმულა (ρ = m/V), მისი ერთეული და "
     "ცდის მიზანი; ხსნის, რისი გაზომვაა საჭირო."],
    ["2. გაზომვების ჩატარება", "3", "სასწორითა და საზომი ცილინდრით სწორად და ზუსტად ზომავს სხეულის მასასა "
     "და მოცულობას; იცავს გაზომვის ტექნიკას."],
    ["3. გამოთვლა და შედეგი", "2", "სწორად იყენებს ფორმულას, აწარმოებს არითმეტიკულად ზუსტ გამოთვლას და "
     "მიუთითებს შედეგის ერთეულს (კგ/მ³ ან გ/სმ³)."],
    ["4. შედეგის ანალიზი და დასკვნა", "2", "შედარებულია მიღებული სიმკვრივე ცხრილურ მონაცემთან; ახსნილია "
     "შესაძლო გადახრის მიზეზი (ცდომილება)."],
    ["5. ანგარიშის გაფორმება და უსაფრთხოება", "1", "ჩანაწერი თანმიმდევრული და სუფთაა; დაცულია ხელსაწყოებთან "
     "მუშაობის უსაფრთხოების წესები."],
]
make_table(doc, rubric_headers, rubric_rows, [6.5, 2.5, 8.0])

total_p = doc.add_paragraph()
total_p.paragraph_format.space_before = Pt(6)
total_p.paragraph_format.space_after = Pt(4)
r = total_p.add_run("სულ: 10 ქულა")
style_run(r, size=11.5, bold=True, color=ACCENT)

add_para(doc, "საბოლოო ნიშანი დგინდება ჯამური ქულის მიხედვით, ცხრილში მე-2 პუნქტში მოცემული 10-ქულიანი სკალის "
              "შესაბამისად.", size=10.5, italic=True, color=RGBColor(0x56, 0x5F, 0x75), space_before=4)

doc.save(r"C:\Users\User\Desktop\my_projects\school\შეფასების სისტემა.docx")
print("saved")
